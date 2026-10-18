"""Generate the sample Task-Tracker.xlsx input (30 fictional PM tasks).

100% local. All data fictional. by Muhammad Bilal.
"""
import datetime as dt

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REF = dt.date(2026, 10, 16)  # "today" for the sample data

TASKS = [
    ("Finalize UAT test plan", "City Lab LIMS Rollout", "Ayesha Khan", "Completed", "High", -12),
    ("Configure instrument interfaces", "City Lab LIMS Rollout", "Bilal Ahmed", "In Progress", "Critical", -3),
    ("Train super-users (batch 1)", "City Lab LIMS Rollout", "Fatima Raza", "In Progress", "High", 2),
    ("Data migration dry run", "City Lab LIMS Rollout", "Hassan Sheikh", "Blocked", "Critical", -5),
    ("Sign-off: deployment checklist", "City Lab LIMS Rollout", "Mahnoor Ali", "Not Started", "High", 10),
    ("Vendor SLA review", "City Lab LIMS Rollout", "Usman Tariq", "Completed", "Medium", -20),
    ("Backup & rollback plan", "City Lab LIMS Rollout", "Ayesha Khan", "In Progress", "High", 4),
    ("Stakeholder demo (steering committee)", "City Lab LIMS Rollout", "Bilal Ahmed", "Not Started", "Medium", 14),
    ("Close defect log (round 2)", "City Lab LIMS Rollout", "Fatima Raza", "In Progress", "Critical", -1),
    ("Go-live communication pack", "City Lab LIMS Rollout", "Hassan Sheikh", "Not Started", "Medium", 21),
    ("Audit CAPA register cleanup", "QMS Digitization", "Mahnoor Ali", "In Progress", "High", -7),
    ("Digitize SOP approval workflow", "QMS Digitization", "Usman Tariq", "Completed", "High", -15),
    ("Internal audit calendar Q4", "QMS Digitization", "Ayesha Khan", "Completed", "Medium", -9),
    ("Management review slides", "QMS Digitization", "Bilal Ahmed", "In Progress", "Medium", 1),
    ("Supplier scorecard template", "QMS Digitization", "Fatima Raza", "Not Started", "Low", 18),
    ("Close overdue CAPAs (Q2)", "QMS Digitization", "Hassan Sheikh", "Blocked", "Critical", -11),
    ("E-QMS user guide", "QMS Digitization", "Mahnoor Ali", "Not Started", "Medium", 12),
    ("Document control audit", "QMS Digitization", "Usman Tariq", "In Progress", "High", 6),
    ("QA dashboard v2", "QMS Digitization", "Ayesha Khan", "In Progress", "Medium", 3),
    ("Retention policy sign-off", "QMS Digitization", "Bilal Ahmed", "Not Started", "Low", 25),
    ("Reconcile migrated counts", "Data Migration & Reporting", "Fatima Raza", "In Progress", "Critical", -2),
    ("Build KPI trend charts", "Data Migration & Reporting", "Hassan Sheikh", "Not Started", "Medium", 8),
    ("Archive legacy reports", "Data Migration & Reporting", "Mahnoor Ali", "Completed", "Low", -30),
    ("Automate weekly KPI email", "Data Migration & Reporting", "Usman Tariq", "In Progress", "High", 5),
    ("Validate turnaround-time data", "Data Migration & Reporting", "Ayesha Khan", "Blocked", "High", -4),
    ("Reporting module UAT", "Data Migration & Reporting", "Bilal Ahmed", "Not Started", "High", 16),
    ("Handover: reporting runbook", "Data Migration & Reporting", "Fatima Raza", "Not Started", "Medium", 28),
    ("Exception report template", "Data Migration & Reporting", "Hassan Sheikh", "Completed", "Medium", -6),
    ("Monthly compliance pack", "Data Migration & Reporting", "Mahnoor Ali", "In Progress", "Medium", 0),
    ("Clean duplicate records", "Data Migration & Reporting", "Usman Tariq", "In Progress", "High", -8),
]

HDR_FILL = PatternFill("solid", fgColor="0E4C81")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Border(*[Side(style="thin", color="B0B0B0")] * 4)


def main() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tasks"
    headers = ["Task ID", "Title", "Project", "Owner", "Status", "Priority", "Due Date"]
    widths = [12, 42, 26, 16, 13, 10, 13]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = THIN
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
    for i, (title, project, owner, status, priority, off) in enumerate(TASKS, start=1):
        row = [f"T-{i:03d}", title, project, owner, status, priority,
               REF + dt.timedelta(days=off)]
        for col, v in enumerate(row, start=1):
            c = ws.cell(row=i + 1, column=col, value=v)
            c.border = THIN
            c.alignment = Alignment(vertical="center", wrap_text=(col == 2))
            if col == 7:
                c.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(TASKS) + 1}"
    wb.save("Task-Tracker.xlsx")
    print(f"Wrote Task-Tracker.xlsx with {len(TASKS)} tasks.")


if __name__ == "__main__":
    main()
