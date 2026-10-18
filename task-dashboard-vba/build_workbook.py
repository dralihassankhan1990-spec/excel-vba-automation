"""Build the base Task-Dashboard workbook (Task Export sheet, 40 sample tasks).

The VBA macro (TaskDashboard.bas) is NOT embedded here — it is injected and
executed by run_macro.py via the LibreOffice UNO bridge, which saves the final
Task-Dashboard.xlsm with the macro + a generated Dashboard sheet.

100% local. All data fictional. by Muhammad Bilal.
"""
import datetime as dt
import random

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TODAY = dt.date(2026, 10, 16)
random.seed(42)

OWNERS = ["Ayesha Khan", "Bilal Ahmed", "Fatima Raza", "Hassan Sheikh", "Mahnoor Ali", "Usman Tariq"]
PROJECTS = ["City Lab LIMS Rollout", "QMS Digitization", "Data Migration & Reporting"]

# (title, project, owner, status, priority, due-offset-days)
TASKS = [
    ("Map current sample-receiving workflow", "City Lab LIMS Rollout", "Ayesha Khan", "Completed", "High", -40),
    ("Draft BRD for sample registration module", "City Lab LIMS Rollout", "Bilal Ahmed", "Completed", "Critical", -35),
    ("Configure user roles & permissions matrix", "City Lab LIMS Rollout", "Fatima Raza", "In Progress", "High", -6),
    ("UAT test cases for accessioning flow", "City Lab LIMS Rollout", "Hassan Sheikh", "In Progress", "High", 3),
    ("Run UAT round 1 with lab analysts", "City Lab LIMS Rollout", "Mahnoor Ali", "In Progress", "Critical", -2),
    ("Fix UAT defects batch 1 (12 items)", "City Lab LIMS Rollout", "Usman Tariq", "Blocked", "Critical", -9),
    ("Integrate barcode printer with registration", "City Lab LIMS Rollout", "Bilal Ahmed", "Not Started", "Medium", 12),
    ("Prepare training deck for lab staff", "City Lab LIMS Rollout", "Ayesha Khan", "In Progress", "Medium", 5),
    ("Data validation rules for test results", "City Lab LIMS Rollout", "Fatima Raza", "Not Started", "High", 18),
    ("Sign-off: go-live readiness checklist", "City Lab LIMS Rollout", "Hassan Sheikh", "Not Started", "Critical", 25),
    ("Archive legacy paper registers (scan)", "City Lab LIMS Rollout", "Mahnoor Ali", "Completed", "Low", -20),
    ("Set up staging environment access", "City Lab LIMS Rollout", "Usman Tariq", "Completed", "Medium", -45),
    ("Weekly status report to steering committee", "City Lab LIMS Rollout", "Ayesha Khan", "In Progress", "Medium", 2),
    ("Vendor follow-up: instrument interface API", "City Lab LIMS Rollout", "Bilal Ahmed", "Blocked", "High", -4),
    ("Document QMS SOP index (v3)", "QMS Digitization", "Fatima Raza", "Completed", "High", -30),
    ("Digitize corrective-action (CAPA) forms", "QMS Digitization", "Hassan Sheikh", "In Progress", "Critical", -5),
    ("Build audit-findings register", "QMS Digitization", "Mahnoor Ali", "Completed", "High", -18),
    ("Migrate 2024 audit records", "QMS Digitization", "Usman Tariq", "In Progress", "Medium", 8),
    ("Review document control workflow", "QMS Digitization", "Ayesha Khan", "Not Started", "Medium", 15),
    ("Train QA team on e-QMS portal", "QMS Digitization", "Bilal Ahmed", "Not Started", "High", 21),
    ("Internal audit schedule Q4", "QMS Digitization", "Fatima Raza", "Completed", "Medium", -12),
    ("Close overdue CAPA items from Q2", "QMS Digitization", "Hassan Sheikh", "In Progress", "Critical", -11),
    ("Management review pack (Q3)", "QMS Digitization", "Mahnoor Ali", "In Progress", "High", 1),
    ("Calibrate document approval reminders", "QMS Digitization", "Usman Tariq", "Blocked", "Low", -3),
    ("Backup & retention policy for e-QMS", "QMS Digitization", "Ayesha Khan", "Not Started", "Medium", 30),
    ("Supplier evaluation form (digital)", "QMS Digitization", "Bilal Ahmed", "Completed", "Low", -25),
    ("Publish QMS dashboard v1", "QMS Digitization", "Fatima Raza", "In Progress", "High", 6),
    ("Extract legacy sample data (2022-2024)", "Data Migration & Reporting", "Hassan Sheikh", "Completed", "Critical", -50),
    ("Clean duplicate sample IDs", "Data Migration & Reporting", "Mahnoor Ali", "In Progress", "High", -8),
    ("Map legacy fields to new schema", "Data Migration & Reporting", "Usman Tariq", "Completed", "Critical", -33),
    ("Build migration validation scripts", "Data Migration & Reporting", "Ayesha Khan", "In Progress", "High", 4),
    ("Trial load: 10k historical records", "Data Migration & Reporting", "Bilal Ahmed", "Blocked", "High", -1),
    ("Reconcile migrated vs source counts", "Data Migration & Reporting", "Fatima Raza", "Not Started", "Critical", 10),
    ("Design weekly KPI report template", "Data Migration & Reporting", "Hassan Sheikh", "Completed", "Medium", -15),
    ("Automate monthly compliance report", "Data Migration & Reporting", "Mahnoor Ali", "Not Started", "Medium", 20),
    ("Dashboard: turnaround-time trends", "Data Migration & Reporting", "Usman Tariq", "In Progress", "Medium", 7),
    ("Archive pre-2022 paper results", "Data Migration & Reporting", "Ayesha Khan", "Completed", "Low", -28),
    ("User acceptance: reporting module", "Data Migration & Reporting", "Bilal Ahmed", "Not Started", "High", 14),
    ("Handover: migration runbook", "Data Migration & Reporting", "Fatima Raza", "Not Started", "Medium", 28),
    ("Prepare data-quality exception report", "Data Migration & Reporting", "Fatima Raza", "In Progress", "Medium", 9),
]

HDR_FILL = PatternFill("solid", fgColor="0E4C81")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Border(*[Side(style="thin", color="B0B0B0")] * 4)


def main() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Task Export"
    headers = ["Task ID", "Title", "Project", "Owner", "Status", "Priority", "Due Date"]
    widths = [12, 44, 26, 16, 13, 10, 13]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = THIN
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 24

    for i, (title, project, owner, status, priority, off) in enumerate(TASKS, start=1):
        due = TODAY + dt.timedelta(days=off)
        row = [
            f"T-{i:03d}",
            title,
            project,
            owner,
            status,
            priority,
            due,
        ]
        for col, v in enumerate(row, start=1):
            c = ws.cell(row=i + 1, column=col, value=v)
            c.border = THIN
            c.alignment = Alignment(vertical="center", wrap_text=(col == 2))
            if col == 7:
                c.number_format = "yyyy-mm-dd"
        ws.row_dimensions[i + 1].height = 20

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(TASKS) + 1}"
    ws.sheet_properties.tabColor = "0E4C81"
    out = "Task-Dashboard-base.xlsx"
    wb.save(out)
    print(f"Wrote {out} with {len(TASKS)} tasks (source for the macro run).")


if __name__ == "__main__":
    main()
